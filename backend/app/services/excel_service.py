# -*- coding: utf-8 -*-
"""
ExcelService: импорт матричного расписания кафедры из файлов .xlsx / .xls
в таблицу public.nnz_schedule (asyncpg, PostgreSQL).

Реальная структура Excel-файлов расписания:
  Строка 1:     Заголовок (пропускается)
  Строка 2:     Шапка групп (колонки C+)
  Далее:        Дата (col A, появляется один раз за день)
                Слот пары (col B: "1-2", "3-4", "5-6", "7-8")
                Каждый слот = 3 строки: предмет / преподаватель / аудитория
"""

import io
import re
import logging
from datetime import date, timedelta, datetime
from typing import Any, Dict, List, Optional, Set, Tuple

import asyncpg
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    import xlrd  # type: ignore
    _XLRD_AVAILABLE = True
except ImportError:
    _XLRD_AVAILABLE = False

from ..schemas.excel import ExcelImportResult, ExcelRowError
from ..repositories.schedule import ScheduleRepository

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Константы
# ---------------------------------------------------------------------------

# Excel-значения, которые означают «пары нет» — тихо пропускаем
_SKIP_VALUES: Set[str] = {
    "", "-", "_", "сам.работа", "сам. работа", "сам.раб", "самостоятельная работа",
    "сам. раб.", "самоподготовка",
}

# Маппинг строки колонки B → номер пары (1-4)
_SLOT_TO_PAIR: Dict[str, int] = {
    "1-2": 1, "3-4": 2, "5-6": 3, "7-8": 4,
    "1": 1, "2": 2, "3": 3, "4": 4,   # запасные варианты
}

# Маппинг буквы типа занятия → pair_type_id (запасной — для cases без eventtools в кэше)
_TYPE_LETTER_FALLBACK: Dict[str, int] = {
    "л": 1, "п": 2, "с": 3, "кр": 4, "лб": 5, "лаб": 5,
}

# Excel epoch offset
_XLS_EPOCH = date(1899, 12, 30)


# ---------------------------------------------------------------------------
# ExcelService
# ---------------------------------------------------------------------------

class ExcelService:
    """Сервис импорта/экспорта расписания через Excel."""

    def __init__(self, conn: asyncpg.Connection) -> None:
        self.conn = conn
        self._repo = ScheduleRepository(conn)

    # ------------------------------------------------------------------
    # Публичный метод импорта
    # ------------------------------------------------------------------

    async def import_schedule_from_excel(
        self,
        file_bytes: bytes,
        filename: str = "schedule.xlsx",
        sh_var_id: Optional[int] = None,
        dry_run: bool = False,
        cathedra_id: Optional[int] = None,
    ) -> ExcelImportResult:
        """
        Парсит матричную сетку расписания и сохраняет в nnz_schedule.

        Args:
            file_bytes:   Содержимое Excel-файла.
            filename:     Имя файла (для определения формата .xls vs .xlsx).
            sh_var_id:    ID варианта расписания. None — определять по датам.
            dry_run:      True → анализировать без записи в БД.
            cathedra_id:  Кафедра для поля idcathedra (опционально).

        Returns:
            ExcelImportResult с подробной статистикой.
        """
        # ---------- загрузка справочников ----------
        try:
            caches = await self._build_lookup_caches()
            s_year_id: int = caches["s_year_id"]
        except Exception as exc:
            import traceback
            traceback.print_exc()
            logger.error("Ошибка при загрузке кэшей справочников БД: %s", exc)
            return ExcelImportResult(
                success=False,
                message=f"Ошибка при инициализации справочников БД: {exc}",
                total_rows=0, imported_count=0, updated_count=0,
                skipped_count=0, failed_count=1,
                errors=[ExcelRowError(row_index=0, reason=f"Справочники БД: {exc}")],
            )

        errors: List[ExcelRowError] = []
        records: List[Dict[str, Any]] = []
        unmatched_teachers: Set[str] = set()
        unmatched_courses: Set[str] = set()
        unmatched_groups: Set[str] = set()
        total_slots = 0

        # ---------- загрузка книги ----------
        try:
            ws_data = self._load_sheet(file_bytes, filename)
        except Exception as exc:
            return ExcelImportResult(
                success=False,
                message=f"Не удалось открыть файл: {exc}",
                total_rows=0, imported_count=0, updated_count=0,
                skipped_count=0, failed_count=1,
                errors=[ExcelRowError(row_index=0, reason=str(exc))],
            )

        # ws_data — list[list[Any]], индексы [row][col]
        if len(ws_data) < 3:
            return ExcelImportResult(
                success=False,
                message="Файл содержит менее 3 строк — не является корректным расписанием.",
                total_rows=0, imported_count=0, updated_count=0,
                skipped_count=0, failed_count=0,
            )

        # ---------- строка 2: читаем заголовки групп ----------
        header_row = ws_data[1]  # row index 1 (0-based)
        # Колонки A=0, B=1, C=2 и далее — группы
        col_to_group_name: Dict[int, str] = {}
        for col_idx in range(2, len(header_row)):
            raw = self._clean(header_row[col_idx])
            if not raw:
                continue
            group_name = self._extract_group_name(raw)
            if group_name:
                col_to_group_name[col_idx] = group_name

        if not col_to_group_name:
            return ExcelImportResult(
                success=False,
                message="В строке 2 не найдено ни одной учебной группы (ожидаются в колонках C+).",
                total_rows=0, imported_count=0, updated_count=0,
                skipped_count=0, failed_count=0,
            )

        logger.info("Найдено групп в заголовке: %d → %s", len(col_to_group_name),
                    list(col_to_group_name.values())[:10])

        # ---------- основной цикл по строкам ----------
        current_date: Optional[date] = None
        row_idx = 2  # начинаем с третьей строки (0-based index 2)
        total_rows = len(ws_data)

        while row_idx < total_rows:
            row = ws_data[row_idx]

            def _cell(r: List[Any], c: int) -> str:
                return self._clean(r[c] if c < len(r) else None)

            col_a = _cell(row, 0)
            col_b = _cell(row, 1)

            # Проверяем дату в колонке A
            if col_a:
                parsed = self._parse_date(col_a)
                if parsed:
                    current_date = parsed

            # Проверяем, является ли col B слотом пары
            pair_num = self._parse_slot(col_b)
            if pair_num is None:
                row_idx += 1
                continue

            # Следующие 2 строки: строка преподавателей и строка аудиторий
            row_teacher = ws_data[row_idx + 1] if row_idx + 1 < total_rows else []
            row_room    = ws_data[row_idx + 2] if row_idx + 2 < total_rows else []

            excel_row_num = row_idx + 1  # 1-based для логов

            # ---------- по каждой группе в столбцах ----------
            for col_idx, group_name in col_to_group_name.items():
                subj_val  = _cell(row, col_idx)
                teach_val = _cell(row_teacher, col_idx)
                room_val  = _cell(row_room, col_idx)

                total_slots += 1

                # Пустой или «Самоподготовка» — пропускаем
                if subj_val.lower() in _SKIP_VALUES:
                    continue

                # Нет даты — не знаем куда класть
                if current_date is None:
                    errors.append(ExcelRowError(
                        row_index=excel_row_num,
                        reason=f"Не определена дата для пары (колонка A пуста выше)"
                    ))
                    continue

                # --- Группа ---
                gid = caches["groups"].get(group_name.lower())
                if not gid:
                    # попробуем нечёткий поиск по числовой части
                    gid = self._fuzzy_group_lookup(group_name, caches["groups"])
                if not gid:
                    unmatched_groups.add(group_name)
                    continue

                # --- Дисциплина ---
                parsed_subj = self._parse_subject(subj_val)
                if not parsed_subj:
                    errors.append(ExcelRowError(
                        row_index=excel_row_num,
                        reason=f"Не удалось распознать дисциплину: '{subj_val}'"
                    ))
                    continue

                course_key   = parsed_subj["course_key"]
                type_letter  = parsed_subj["type_letter"]
                lesson_num   = parsed_subj["lesson_num"]

                cid = caches["courses"].get(course_key.lower())
                if not cid:
                    # попробуем без суффиксов (ДМ-42,52,62 → ДМ-42)
                    cid = self._fuzzy_course_lookup(course_key, caches["courses"])
                if not cid:
                    unmatched_courses.add(course_key)
                    continue

                # --- Тип занятия ---
                pair_type_id = self._resolve_type(type_letter, caches["types"])

                # --- Преподаватель ---
                mid: Optional[int] = None
                if teach_val and teach_val.lower() not in _SKIP_VALUES:
                    teacher_key = self._parse_teacher(teach_val)
                    if teacher_key:
                        mid = self._lookup_teacher(teacher_key, caches)
                    if not mid:
                        unmatched_teachers.add(teach_val)
                        continue
                else:
                    # Преподаватель не указан — пропускаем
                    continue

                # --- Аудитория (опционально) ---
                rid_list: List[int] = []
                if room_val and room_val.lower() not in _SKIP_VALUES:
                    rid = caches["rooms"].get(room_val.lower())
                    if rid:
                        rid_list.append(rid)

                # --- sh_var_id: из параметра или по дате ---
                resolved_var_id = sh_var_id
                if resolved_var_id is None:
                    try:
                        resolved_var_id = await self._repo.get_sh_var_id_by_date(
                            current_date, s_year_id
                        )
                    except Exception as exc:
                        logger.warning("Не удалось определить sh_var_id для %s: %s", current_date, exc)
                if not resolved_var_id:
                    errors.append(ExcelRowError(
                        row_index=excel_row_num,
                        reason=f"Не найден вариант расписания для даты {current_date}"
                    ))
                    continue

                # --- period lid ---
                try:
                    period_lid = await self._repo.get_period_lid(s_year_id, pair_num)
                except Exception as exc:
                    logger.warning("Не удалось определить period lid для пары %d: %s", pair_num, exc)
                    period_lid = pair_num  # fallback

                # --- День недели (1=ПН ... 7=ВС) ---
                day_of_week = current_date.isoweekday()  # 1-7, Mon=1

                records.append({
                    "day_of_week": day_of_week,
                    "period": period_lid,
                    "gid": gid,
                    "cid": cid,
                    "pair_type_id": pair_type_id,
                    "teacher_mid": [mid],
                    "rid": rid_list,
                    "sh_var_id": resolved_var_id,
                    "idcathedra": cathedra_id,
                    "lesson_num": lesson_num,
                    # debug info (не вставляем в БД)
                    "_date": current_date,
                    "_pair_num": pair_num,
                    "_group": group_name,
                    "_excel_row": excel_row_num,
                })

            # Перескакиваем через все 3 строки слота
            row_idx += 3

        # ---------- итоговая статистика ----------
        skipped = total_slots - len(records) - len(errors)
        logger.info(
            "Парсинг завершён: total_slots=%d valid=%d errors=%d skipped=%d",
            total_slots, len(records), len(errors), max(0, skipped)
        )
        logger.info(
            "Ненайдено: teachers=%d courses=%d groups=%d",
            len(unmatched_teachers), len(unmatched_courses), len(unmatched_groups)
        )

        # ---------- запись в БД ----------
        inserted_count = 0
        updated_count = 0
        created_ids: List[int] = []

        if not dry_run and records:
            try:
                async with self.conn.transaction():
                    for item in records:
                        existing = await self.conn.fetchval(
                            """
                            SELECT sheid FROM public.nnz_schedule
                            WHERE sh_var_id = $1 AND day_of_week = $2
                              AND period = $3 AND gid = $4
                            LIMIT 1;
                            """,
                            item["sh_var_id"],
                            item["day_of_week"],
                            item["period"],
                            item["gid"],
                        )
                        if existing:
                            await self.conn.execute(
                                """
                                UPDATE public.nnz_schedule
                                SET cid=$1, pair_type_id=$2, teacher_mid=$3,
                                    rid=$4, idcathedra=$5, lesson_num=$6
                                WHERE sheid=$7;
                                """,
                                item["cid"], item["pair_type_id"],
                                item["teacher_mid"], item["rid"],
                                item["idcathedra"], item["lesson_num"],
                                existing,
                            )
                            updated_count += 1
                            created_ids.append(existing)
                        else:
                            new_id = await self.conn.fetchval(
                                """
                                INSERT INTO public.nnz_schedule (
                                    day_of_week, period, gid, cid, pair_type_id,
                                    teacher_mid, rid, sh_var_id, idcathedra, lesson_num
                                ) VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10)
                                RETURNING sheid;
                                """,
                                item["day_of_week"], item["period"],
                                item["gid"], item["cid"], item["pair_type_id"],
                                item["teacher_mid"], item["rid"],
                                item["sh_var_id"], item["idcathedra"],
                                item["lesson_num"],
                            )
                            if new_id:
                                inserted_count += 1
                                created_ids.append(new_id)
            except Exception as exc:
                import traceback
                traceback.print_exc()
                return ExcelImportResult(
                    success=False,
                    message=f"Ошибка при записи в базу данных: {exc}",
                    total_rows=total_slots,
                    imported_count=0, updated_count=0,
                    skipped_count=total_slots,
                    failed_count=1,
                    errors=errors[:200],
                    unmatched_teachers=sorted(unmatched_teachers),
                    unmatched_courses=sorted(unmatched_courses),
                    unmatched_groups=sorted(unmatched_groups),
                )
        else:
            # dry_run: считаем как «потенциально импортируемые»
            inserted_count = len(records)

        total_failed = len(errors)
        total_skipped = total_slots - len(records) - total_failed
        success_flag = (inserted_count + updated_count) > 0 or dry_run

        msg_parts = [
            f"Всего слотов: {total_slots}.",
            f"{'(dry_run) ' if dry_run else ''}Добавлено: {inserted_count}, обновлено: {updated_count}.",
            f"Пропущено: {max(0, total_skipped)}, ошибок: {total_failed}.",
        ]
        if unmatched_teachers:
            msg_parts.append(f"Не найдено преподавателей: {len(unmatched_teachers)}.")
        if unmatched_courses:
            msg_parts.append(f"Не найдено дисциплин: {len(unmatched_courses)}.")
        if unmatched_groups:
            msg_parts.append(f"Не найдено групп: {len(unmatched_groups)}.")

        return ExcelImportResult(
            success=success_flag,
            message=" ".join(msg_parts),
            total_rows=total_slots,
            imported_count=inserted_count,
            updated_count=updated_count,
            skipped_count=max(0, total_skipped),
            failed_count=total_failed,
            stats={
                "inserted": inserted_count,
                "updated": updated_count,
                "skipped": max(0, total_skipped),
                "dry_run": dry_run,
                "valid_records": len(records),
            },
            errors=errors[:200],
            created_ids=created_ids,
            unmatched_teachers=sorted(unmatched_teachers),
            unmatched_courses=sorted(unmatched_courses),
            unmatched_groups=sorted(unmatched_groups),
        )

    # ------------------------------------------------------------------
    # Загрузка книги
    # ------------------------------------------------------------------

    def _load_sheet(self, file_bytes: bytes, filename: str) -> List[List[Any]]:
        """
        Читает Excel-файл и возвращает данные как list[list[Any]].
        Для .xlsx использует openpyxl (с распаковкой merged cells).
        Для .xls использует xlrd.
        """
        fname_lower = filename.lower()

        if fname_lower.endswith(".xls") and not fname_lower.endswith(".xlsx"):
            return self._load_xls(file_bytes)
        else:
            return self._load_xlsx(file_bytes)

    def _load_xlsx(self, file_bytes: bytes) -> List[List[Any]]:
        """Загружает .xlsx через openpyxl, разворачивает объединённые ячейки."""
        wb = openpyxl.load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
        ws = wb.active

        # Распаковка объединённых ячеек: копируем значение из верхней левой
        # ячейки во все остальные ячейки диапазона
        merged_ranges = list(ws.merged_cells.ranges)
        for merge_range in merged_ranges:
            top_left_value = ws.cell(merge_range.min_row, merge_range.min_col).value
            ws.unmerge_cells(str(merge_range))
            for row in range(merge_range.min_row, merge_range.max_row + 1):
                for col in range(merge_range.min_col, merge_range.max_col + 1):
                    ws.cell(row, col).value = top_left_value

        data: List[List[Any]] = []
        for row in ws.iter_rows(values_only=True):
            data.append(list(row))
        return data

    def _load_xls(self, file_bytes: bytes) -> List[List[Any]]:
        """Загружает .xls через xlrd (Excel 97-2003)."""
        if not _XLRD_AVAILABLE:
            raise ImportError(
                "Для чтения файлов .xls необходима библиотека xlrd. "
                "Установите её: pip install xlrd==1.2.0"
            )
        book = xlrd.open_workbook(file_contents=file_bytes)
        sheet = book.sheet_by_index(0)

        data: List[List[Any]] = []
        for row_idx in range(sheet.nrows):
            row_data: List[Any] = []
            for col_idx in range(sheet.ncols):
                cell = sheet.cell(row_idx, col_idx)
                # xlrd типы: 0=empty, 1=text, 2=number, 3=date, 4=bool, 5=error
                if cell.ctype == 3:  # date serial
                    try:
                        dt_tuple = xlrd.xldate_as_tuple(cell.value, book.datemode)
                        row_data.append(datetime(*dt_tuple[:6]).date() if dt_tuple[0] else None)
                    except Exception:
                        row_data.append(cell.value)
                else:
                    row_data.append(cell.value)
            data.append(row_data)
        return data

    # ------------------------------------------------------------------
    # Кэш справочников
    # ------------------------------------------------------------------

    async def _build_lookup_caches(self) -> Dict[str, Any]:
        """
        Загружает все справочники БД в память одним запросом каждый.
        Возвращает словарь с именованными картами для быстрого поиска.
        """
        # --- Учебный год ---
        year_row = await self.conn.fetchrow(
            "SELECT school_year FROM public.studyyears ORDER BY school_year DESC LIMIT 1;"
        )
        s_year_id: int = year_row["school_year"] if year_row else 35

        # --- Группы: {name_lower: gid} ---
        grp_rows = await self.conn.fetch(
            "SELECT gid, LOWER(TRIM(name)) AS key FROM public.groupname;"
        )
        groups_map: Dict[str, int] = {r["key"]: r["gid"] for r in grp_rows}

        # --- Дисциплины: {alias_lower: cid, title_lower: cid} ---
        crs_rows = await self.conn.fetch(
            "SELECT cid, LOWER(TRIM(COALESCE(alias,''))) AS alias, "
            "LOWER(TRIM(COALESCE(title,''))) AS title FROM public.courses;"
        )
        courses_map: Dict[str, int] = {}
        for r in crs_rows:
            if r["alias"]:
                courses_map[r["alias"]] = r["cid"]
            if r["title"]:
                courses_map[r["title"]] = r["cid"]

        # --- Типы занятий: {alias_lower: typeid} (в таблице eventtools колонка alias, а не name!) ---
        type_rows = await self.conn.fetch(
            "SELECT typeid, LOWER(TRIM(COALESCE(alias, ''))) AS alias FROM public.eventtools;"
        )
        types_map: Dict[str, int] = {}
        for r in type_rows:
            if r["alias"]:
                types_map[r["alias"]] = r["typeid"]

        # --- Аудитории: {short_name_lower: rid} (колонка short_name, а не name!) ---
        room_rows = await self.conn.fetch(
            "SELECT rid, LOWER(TRIM(short_name)) AS key FROM public.rooms WHERE short_name IS NOT NULL;"
        )
        rooms_map: Dict[str, int] = {r["key"]: r["rid"] for r in room_rows}

        # --- Преподаватели кафедры: {(lastname_nyo, init1, init2): mid} ---
        # Только преподаватели кафедры (cathedra_personnel, id_pmk IN (1, 2))
        teacher_rows = await self.conn.fetch(
            """
            SELECT DISTINCT p.mid,
                   LOWER(TRIM(p.lastname))   AS lastname,
                   LOWER(LEFT(TRIM(p.firstname), 1))   AS init1,
                   LOWER(LEFT(TRIM(COALESCE(p.patronymic, '')), 1)) AS init2
            FROM public.people p
            JOIN public.cathedra_personnel cp ON cp.mid = p.mid
            WHERE cp.id_pmk IN (1, 2)
              AND p.lastname IS NOT NULL
              AND p.firstname IS NOT NULL;
            """
        )
        teachers_map: Dict[Tuple[str, str, str], int] = {}
        teachers_map_no_patronymic: Dict[Tuple[str, str], int] = {}
        for r in teacher_rows:
            last = self._normalize_yo(r["lastname"])
            i1 = r["init1"] or ""
            i2 = r["init2"] or ""
            key3 = (last, i1, i2)
            key2 = (last, i1)
            if key3 not in teachers_map:
                teachers_map[key3] = r["mid"]
            if key2 not in teachers_map_no_patronymic:
                teachers_map_no_patronymic[key2] = r["mid"]

        # --- Варианты расписания: sh_var_name (колонка sh_var_name, а не name!) ---
        variants = await self.conn.fetch(
            "SELECT sh_var_id, sh_var_name, s_year_id FROM public.nnz_schedule_variants;"
        )
        variants_map = {r["sh_var_id"]: r["sh_var_name"] for r in variants}

        return {
            "s_year_id": s_year_id,
            "groups": groups_map,
            "courses": courses_map,
            "types": types_map,
            "rooms": rooms_map,
            "teachers": teachers_map,
            "teachers2": teachers_map_no_patronymic,
            "variants": variants_map,
        }

    # ------------------------------------------------------------------
    # Вспомогательные методы
    # ------------------------------------------------------------------

    @staticmethod
    def _clean(val: Any) -> str:
        """Приводит ячейку к строке, убирает лишние пробелы."""
        if val is None:
            return ""
        s = str(val).strip()
        # убираем внутренние множественные пробелы
        s = re.sub(r"\s{2,}", " ", s)
        return s

    @staticmethod
    def _normalize_yo(s: str) -> str:
        """Заменяет ё→е для нормализованного сравнения."""
        return s.replace("ё", "е").replace("Ё", "Е")

    @staticmethod
    def _parse_date(val: Any) -> Optional[date]:
        """
        Разбирает дату из ячейки Excel.
        Форматы: "31.08.2026", "31.08.26", datetime, date, float (serial).
        """
        if val is None:
            return None
        if isinstance(val, (datetime,)):
            return val.date()
        if isinstance(val, date):
            return val
        if isinstance(val, float) and val > 1:
            # Excel serial date
            try:
                return _XLS_EPOCH + timedelta(days=int(val))
            except Exception:
                return None
        s = str(val).strip()
        for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _parse_slot(val: str) -> Optional[int]:
        """
        Распознаёт строку слота пары из колонки B.
        "1-2" → 1, "3-4" → 2, "5-6" → 3, "7-8" → 4.
        Возвращает None если строка не является слотом.
        """
        v = val.strip()
        if v in _SLOT_TO_PAIR:
            return _SLOT_TO_PAIR[v]
        # Может быть "1 - 2", "1–2" и т.п.
        normalized = re.sub(r"\s*[-–—]\s*", "-", v)
        return _SLOT_TO_PAIR.get(normalized)

    @staticmethod
    def _extract_group_name(raw: str) -> Optional[str]:
        """
        Извлекает шифр учебной группы из заголовка столбца.
        Приоритеты:
        1. Четырёхзначный номер: 1180, 1181
        2. Пятизначный: 10001
        3. Буквенно-цифровой: 41ПП, 41/1ПП, 111А
        """
        if not raw:
            return None
        # Четырёх-/пятизначное число
        m = re.search(r"\b(\d{4,5})\b", raw)
        if m:
            return m.group(1)
        # Буквенно-цифровое с дробью (41/1ПП) — с начала строки или слова
        m = re.search(r"\b(\d{1,3}/\d+[А-ЯЁа-яёA-Z]+)\b", raw)
        if m:
            return m.group(1)
        # Буквенно-цифровое без дроби (41ПП, 111А)
        m = re.search(r"\b(\d{1,3}[А-ЯЁа-яёA-Z]{1,3})\b", raw)
        if m:
            return m.group(1)
        # Fallback: первый непустой токен
        tokens = raw.split()
        return tokens[0] if tokens else None

    @staticmethod
    def _parse_subject(text: str) -> Optional[Dict[str, Any]]:
        """
        Разбирает строку дисциплины.
        Примеры:
          "ДС-191_Л1 Т: 1"      → {course_key: "ДС-191", type_letter: "л", lesson_num: 1}
          "ДМ-42,52,62_Л1 Т: 1" → {course_key: "ДМ-42", type_letter: "л", lesson_num: 1}
          "ДМ-191_П2 Т: 3"      → {course_key: "ДМ-191", type_letter: "п", lesson_num: 3}
        """
        text = text.strip()
        if not text or text.lower() in _SKIP_VALUES:
            return None

        # Извлекаем lesson_num из "Т: N" или "Т:N"
        lesson_num: Optional[int] = None
        m_lesson = re.search(r"[ТT]:\s*(\d+)", text, re.IGNORECASE)
        if m_lesson:
            lesson_num = int(m_lesson.group(1))

        # Извлекаем тип занятия — одна или две буквы после "_"
        type_letter = "л"  # лекция по умолчанию
        m_type = re.search(r"_([А-ЯЁA-Zа-яёa-z]{1,3})\d*", text)
        if m_type:
            type_letter = m_type.group(1).lower()

        # Извлекаем код курса — всё до "_" (или до пробела если нет "_")
        course_key = text
        if "_" in text:
            course_key = text.split("_")[0].strip()
        elif " " in text:
            course_key = text.split()[0].strip()

        # Очищаем запятые (ДМ-42,52,62 → ДМ-42)
        if "," in course_key:
            course_key = course_key.split(",")[0].strip()

        if not course_key:
            return None

        return {
            "course_key": course_key,
            "type_letter": type_letter,
            "lesson_num": lesson_num,
        }

    def _parse_teacher(self, text: str) -> Optional[Tuple[str, str, str]]:
        """
        Разбирает ФИО преподавателя из ячейки Excel.
        Формат: "ФЁДОРОВА О.Н." → ("федорова", "о", "н")
        Нормализует ё→е, приводит к нижнему регистру.
        """
        text = text.strip()
        if not text or text.lower() in _SKIP_VALUES:
            return None

        # Ищем паттерн: ФАМИЛИЯ И.О. или Фамилия И. О.
        m = re.search(
            r"([А-ЯЁа-яё][А-ЯЁа-яё\-]+)\s+([А-ЯЁа-яёA-Z])\.?\s*([А-ЯЁа-яёA-Z])\.?",
            text,
        )
        if not m:
            return None

        last  = self._normalize_yo(m.group(1).lower())
        init1 = self._normalize_yo(m.group(2).lower())
        init2 = self._normalize_yo(m.group(3).lower())
        return (last, init1, init2)

    def _resolve_type(self, type_letter: str, types_map: Dict[str, int]) -> int:
        """Определяет pair_type_id по букве типа (Л/П/С/КР/ЛБ)."""
        if not type_letter:
            return 1
        tl = type_letter.lower()
        # Поиск в кэше БД по alias (л, п, с, кр, лб...)
        if tl in types_map:
            return types_map[tl]
        # Fallback
        return _TYPE_LETTER_FALLBACK.get(tl, 1)

    @staticmethod
    def _fuzzy_group_lookup(name: str, groups_map: Dict[str, int]) -> Optional[int]:
        """Нечёткий поиск группы: ищет числовой префикс в ключах карты."""
        m = re.search(r"\d{3,5}", name)
        if not m:
            return None
        digits = m.group(0)
        # Точное совпадение числа
        if digits in groups_map:
            return groups_map[digits]
        # Совпадение начала ключа
        for key, gid in groups_map.items():
            if key.startswith(digits) or digits.startswith(key):
                return gid
        return None

    @staticmethod
    def _fuzzy_course_lookup(course_key: str, courses_map: Dict[str, int]) -> Optional[int]:
        """Нечёткий поиск дисциплины по началу шифра."""
        key_lower = course_key.lower().strip()
        # Точное
        if key_lower in courses_map:
            return courses_map[key_lower]
        # Поиск ключа, начинающегося с нашего значения
        for db_key, cid in courses_map.items():
            if db_key.startswith(key_lower) or key_lower.startswith(db_key):
                return cid
        return None

    # ------------------------------------------------------------------
    # Поиск преподавателя с fallback
    # ------------------------------------------------------------------

    def _lookup_teacher(
        self, teacher_key: Tuple[str, str, str], caches: Dict[str, Any]
    ) -> Optional[int]:
        """
        Ищет преподавателя сначала по 3-tuple (last, i1, i2),
        затем по 2-tuple (last, i1) если отчество не дало совпадения.
        """
        mid = caches["teachers"].get(teacher_key)
        if mid:
            return mid
        # Без отчества
        mid = caches["teachers2"].get((teacher_key[0], teacher_key[1]))
        return mid

    # ------------------------------------------------------------------
    # Экспорт расписания в Excel (оригинальный метод — сохранён)
    # ------------------------------------------------------------------

    async def export_schedule_to_excel(
        self,
        group_id: Optional[int] = None,
        teacher_id: Optional[int] = None,
        sh_var_id: Optional[int] = 1,
    ) -> io.BytesIO:
        """Формирование красиво оформленного отчета расписания в Excel."""
        conditions = ["s.sh_var_id = $1"]
        params: List[Any] = [sh_var_id]

        if group_id:
            params.append(group_id)
            conditions.append(f"s.gid = ${len(params)}")

        if teacher_id:
            params.append(teacher_id)
            conditions.append(f"${len(params)} = ANY(s.teacher_mid)")

        sql = f"""
            SELECT
                s.day_of_week,
                s.period,
                COALESCE(g.name, '-') AS group_name,
                COALESCE(c.alias, c.title, '-') AS course_name,
                COALESCE(e.alias, '-') AS event_type,
                (
                    SELECT TRIM(CONCAT_WS(' ', p.lastname, p.firstname, p.patronymic))
                    FROM people p WHERE p.mid = ANY(s.teacher_mid) LIMIT 1
                ) AS teacher_name,
                (
                    SELECT r.short_name
                    FROM rooms r WHERE r.rid = ANY(s.rid) LIMIT 1
                ) AS room_name
            FROM public.nnz_schedule s
            LEFT JOIN groupname g ON g.gid = s.gid
            LEFT JOIN courses c ON c.cid = s.cid
            LEFT JOIN eventtools e ON e.typeid = s.pair_type_id
            WHERE {' AND '.join(conditions)}
            ORDER BY s.day_of_week, s.period;
        """
        rows = await self.conn.fetch(sql, *params)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Расписание занятий"

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Arial", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        headers = ["День недели", "Пара", "Группа", "Дисциплина", "Тип занятия", "Преподаватель", "Аудитория"]
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        day_names = {1: "ПН", 2: "ВТ", 3: "СР", 4: "ЧТ", 5: "ПТ", 6: "СБ", 7: "ВС"}
        for row_data in rows:
            ws.append([
                day_names.get(row_data["day_of_week"], str(row_data["day_of_week"])),
                row_data["period"],
                row_data["group_name"],
                row_data["course_name"],
                row_data["event_type"],
                row_data["teacher_name"] or "-",
                row_data["room_name"] or "-",
            ])

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=7):
            for cell in row:
                cell.font = regular_font
                cell.border = thin_border
                if cell.column in (1, 2, 7):
                    cell.alignment = center_align

        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
